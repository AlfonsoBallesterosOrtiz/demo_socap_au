import io
from datetime import date
import streamlit as st

# Importaciones para exportación de documentos
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- CONFIGURACIÓN DE PÁGINA ---
st.set_page_config(
    page_title="Calculadora Regulatoria SOCAP - Anexo U",
    page_icon="🏦",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Inyección de estilos CSS para pulido visual
st.markdown("""
<style>
    .main { background-color: #F8F9FA; }
    .metric-container {
        background-color: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-radius: 8px;
        padding: 15px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.05);
    }
    .badge-card {
        padding: 24px;
        border-radius: 12px;
        text-align: center;
        color: white;
        margin-bottom: 20px;
    }
    .badge-cat-a { background: linear-gradient(135deg, #2E7D32, #1B5E20); border: 2px solid #1B5E20; }
    .badge-cat-b { background: linear-gradient(135deg, #1565C0, #0D47A1); border: 2px solid #0D47A1; }
    .badge-cat-c { background: linear-gradient(135deg, #EF6C00, #E65100); border: 2px solid #BF360C; }
    .badge-cat-d { background: linear-gradient(135deg, #C62828, #B71C1C); border: 2px solid #7F0000; }
    .bp-card {
        background-color: #F0F7FD;
        border-left: 5px solid #0275D8;
        padding: 16px;
        border-radius: 6px;
        margin-bottom: 12px;
    }
</style>
""", unsafe_allow_html=True)

# --- ENCABEZADO PRINCIPAL ---
st.title("🏦 Calculadora Regulatoria Anexo U (SOCAP Nivel Básico)")
st.caption("Cómputo en memoria de Nivel de Capitalización (NICAP), Categoría Regulatoria y Mejores Prácticas Nivel I")

# ==========================================
# DATOS GENERALES
# ==========================================
with st.container():
    st.subheader("📋 Datos Generales de la Sociedad")
    col_dg1, col_dg2, col_dg3 = st.columns(3)
    with col_dg1:
        sociedad = st.text_input("Nombre de la Cooperativa", value="SOCAP La Esperanza S.C. de A.P. de R.L. de C.V.")
    with col_dg2:
        fecha_corte = st.date_input("Fecha de Corte", value=date.today())
    with col_dg3:
        tipo_cierre = st.selectbox("Tipo de Cierre", ["Trimestral", "Semestral", "Anual"])

st.markdown("---")

# ==========================================
# CAPTURA DE DATOS EN DOS COLUMNAS
# ==========================================
col_captura_izq, col_captura_der = st.columns(2)

with col_captura_izq:
    st.subheader("📊 Bloque 1: Cartera de Crédito (Anexo U)")
    st.caption("Cifras en miles de pesos mexicanos (MXN)")
    c1 = st.number_input(
        "(1) Cartera de crédito vigente", 
        min_value=0.0, value=12500.0, step=100.0, format="%.2f",
        help="Saldo de capital de préstamos al corriente sin pagos vencidos."
    )
    c2 = st.number_input(
        "(2) Cartera de crédito vencida", 
        min_value=0.0, value=450.0, step=50.0, format="%.2f",
        help="Saldo de capital de préstamos con amortizaciones no liquidadas en la fecha pactada."
    )
    c3 = st.number_input(
        "(3) Estimación preventiva para riesgos crediticios", 
        min_value=0.0, value=320.0, step=50.0, format="%.2f",
        help="Reserva preventiva constituida en resultados para absorber pérdidas crediticias esperadas."
    )

    st.subheader("🏛️ Bloque 2: Capital Contable y Deducciones")
    c6 = st.number_input(
        "(6) Capital Contable total", 
        min_value=0.0, value=2200.0, step=100.0, format="%.2f",
        help="Suma total de aportaciones ordinarias, reservas y utilidades de la cooperativa."
    )
    c7 = st.number_input(
        "(7) Certificados excedentes que NO cumplen características", 
        min_value=0.0, value=0.0, step=10.0, format="%.2f",
        help="Deducción de certificados voluntarios que no tienen permanencia mínima de 5 años según la norma."
    )
    c8 = st.number_input(
        "(8) Financiamiento para adquisición de partes sociales propias", 
        min_value=0.0, value=0.0, step=10.0, format="%.2f",
        help="Deducción de préstamos otorgados por la cooperativa a sus socios para comprar su propia aportación social."
    )

with col_captura_der:
    st.subheader("⚖️ Bloque 6: Validación de Cuadre Contable")
    st.caption("Verificación de la ecuación patrimonial: Activo = Pasivo + Capital")
    c_activo = st.number_input(
        "Total Activo", 
        min_value=0.0, value=15200.0, step=100.0, format="%.2f",
        help="Total registrado en el Balance General."
    )
    c_pasivo = st.number_input(
        "Total Pasivo", 
        min_value=0.0, value=13000.0, step=100.0, format="%.2f",
        help="Total de depósitos y pasivos exigibles registrados en el Balance General."
    )

    st.subheader("🛡️ Cumplimiento y Antecedentes (LRASCAP Art. 15)")
    cumple_anexo = st.checkbox(
        "La información financiera cumple con las normas del Anexo T y formatos del Anexo U",
        value=True,
        help="Marcar si las balanzas y carteras fueron integradas con la metodología regulatoria aplicable."
    )
    es_segunda_vez_c = st.checkbox(
        "¿Fue Categoría C también en el semestre inmediato anterior?",
        value=False,
        help="Si resulta en Categoría C este semestre y se marca esta casilla, causará degradación obligatoria a Categoría D (Art. 15 LRASCAP)."
    )

# ==========================================
# CÁLCULOS REGULATORIOS EN MEMORIA
# ==========================================
# Bloque 1
c4_cartera_neta = c1 + c2 - c3

# Bloque 2
c9_capital_neto = c6 - c7 - c8

# Bloque 3: Requerimiento y NICAP
c5_req_capital = c4_cartera_neta * 0.08
if c5_req_capital <= 0:
    nicap = 999.99
else:
    nicap = round((c9_capital_neto / c5_req_capital) * 100, 2)

# Bloque 4: Clasificación en Categoría
cat_nombre = ""
cat_desc = ""
cat_css = ""
cat_motivo_degradacion = ""

if nicap >= 150.0 and cumple_anexo:
    cat_nombre = "Categoría A"
    cat_desc = "Riesgo Bajo"
    cat_css = "badge-cat-a"
elif 100.0 <= nicap < 150.0 and cumple_anexo:
    cat_nombre = "Categoría B"
    cat_desc = "Riesgo Moderadamente Bajo"
    cat_css = "badge-cat-b"
elif (50.0 <= nicap < 100.0) or (nicap >= 100.0 and not cumple_anexo):
    if es_segunda_vez_c:
        cat_nombre = "Categoría D"
        cat_desc = "Causal de Disolución y Liquidación"
        cat_css = "badge-cat-d"
        cat_motivo_degradacion = "Segunda clasificación consecutiva en Categoría C (Art. 15, fracción III LRASCAP)."
    else:
        cat_nombre = "Categoría C"
        cat_desc = "Riesgo de Insolvencia"
        cat_css = "badge-cat-c"
else:
    cat_nombre = "Categoría D"
    cat_desc = "Causal de Disolución y Liquidación"
    cat_css = "badge-cat-d"

# Cuadre contable
diferencia_balance = round(c_activo - (c_pasivo + c6), 2)
cuadra_balance = abs(diferencia_balance) <= 1.0

st.markdown("---")

# ==========================================
# RESULTADOS EJECUTIVOS Y CLASIFICACIÓN
# ==========================================
st.subheader("🎯 Resultado Oficial de Capitalización")

# Cuadre Contable Banner
if cuadra_balance:
    st.success("✅ Validación Contable Correcta: El Balance General cuadra perfectamente (Activo = Pasivo + Capital Contable).")
else:
    st.error(f"❌ Balance Descuadrado: Revisa las cifras antes de emitir. Activo − (Pasivo + Capital) = ${diferencia_balance:,.2f} miles de pesos.")

# Tarjeta Gigante de Categoría Regulatoria
card_html = f"""
<div class="badge-card {cat_css}">
    <h3 style="margin: 0; font-size: 1.3rem; text-transform: uppercase; letter-spacing: 1px;">Clasificación Regulatoria Oficial</h3>
    <h1 style="margin: 8px 0; font-size: 3.2rem; font-weight: 800;">{cat_nombre}</h1>
    <h3 style="margin: 0; font-size: 1.4rem; font-weight: 400;">{cat_desc}</h3>
    {f'<p style="margin-top: 10px; font-weight: bold; background: rgba(0,0,0,0.25); padding: 6px 12px; border-radius: 4px;">⚠️ {cat_motivo_degradacion}</p>' if cat_motivo_degradacion else ''}
</div>
"""
st.markdown(card_html, unsafe_allow_html=True)

# Métricas Clave
m_col1, m_col2, m_col3, m_col4 = st.columns(4)
with m_col1:
    st.metric(label="(4) Cartera de Crédito Neta", value=f"${c4_cartera_neta:,.2f}", help="(1) Vigente + (2) Vencida − (3) Estimación")
with m_col2:
    st.metric(label="(9) Capital Neto", value=f"${c9_capital_neto:,.2f}", help="(6) Capital Contable − (7) Deducción Cert. − (8) Deducción Finan.")
with m_col3:
    st.metric(label="(5) Req. Capitalización (8%)", value=f"${c5_req_capital:,.2f}", help="8% del Total de Cartera Neta")
with m_col4:
    st.metric(label="NICAP Obtenido", value=f"{nicap:,.2f}%", help="[(9) Capital Neto / (5) Requerimiento] × 100")

# ==========================================
# BLOQUE 5: MEJORES PRÁCTICAS NIVEL I
# ==========================================
st.markdown("---")
st.markdown("""
<div class="bp-card">
    <h4 style="color: #0275D8; margin: 0 0 5px 0;">💎 Módulo de Mejores Prácticas — Nivel I</h4>
    <p style="margin: 0; font-size: 0.95rem; color: #4A5568;">
        Indicadores prudenciales complementarios para monitoreo interno del Comité de Riesgos y Consejo de Administración.
    </p>
</div>
""", unsafe_allow_html=True)

bp_col1, bp_col2, bp_col3 = st.columns(3)

with bp_col1:
    st.markdown("**1. Coeficiente de Liquidez (≤30 días)**")
    act_liq = st.number_input("Activos líquidos ≤ 30 días", min_value=0.0, value=1500.0, step=100.0, format="%.2f")
    pas_exig = st.number_input("Pasivos exigibles ≤ 30 días", min_value=0.0, value=8000.0, step=100.0, format="%.2f")
    
    if pas_exig > 0:
        ratio_liq = round((act_liq / pas_exig) * 100, 2)
    else:
        ratio_liq = 100.0

    st.metric("Coeficiente de Liquidez", f"{ratio_liq:.2f}%")
    if ratio_liq < 10.0:
        st.error("⚠️ Riesgo de liquidez: se recomienda restringir el otorgamiento de nuevos créditos a plazo largo hasta restituir el colchón de disponibilidades por encima del 10%.")
    else:
        st.success("✅ Posición de tesorería solvente (cumple umbral prudencial ≥10%).")

with bp_col2:
    st.markdown("**2. Cobertura de Bienes Adjudicados**")
    val_adj = st.number_input("Valor total bienes adjudicados", min_value=0.0, value=650.0, step=50.0, format="%.2f")
    res_adj = st.number_input("Reserva registrada sobre adjudicados", min_value=0.0, value=260.0, step=20.0, format="%.2f")
    
    if val_adj > 0:
        cob_adj = round((res_adj / val_adj) * 100, 2)
    else:
        cob_adj = 100.0

    st.metric("% Cobertura Reserva (referencial)", f"{cob_adj:.2f}%")
    st.info("ℹ️ Cobertura analítica global. La versión enterprise audita la antigüedad individual mueble/inmueble según Art. 43.")

with bp_col3:
    st.markdown("**3. Monitoreo Límite Nivel Básico (UDIS)**")
    act_udis = st.number_input("Activos totales expresados en UDIS", min_value=0.0, value=2100000.0, step=50000.0, format="%.2f")
    
    st.metric("Consumo de Límite", f"{(act_udis / 2500000.0) * 100:.1f}%")
    if act_udis >= 2500000.0:
        st.error("🚨 Se superó el límite de Nivel Básico (2,500,000 UDIS). La entidad debe operar obligatoriamente bajo el régimen de Nivel I.")
    elif act_udis >= 2300000.0:
        st.warning("⚠️ Aproximándose al límite de 2,500,000 UDIS de Nivel Básico (≥92%). Se recomienda que el Consejo programe la transición a Nivel I.")
    else:
        st.success("✅ Crecimiento dentro de los límites autorizados para Nivel Básico.")

# ==========================================
# GENERACIÓN DE DOCUMENTOS (MEMORIA)
# ==========================================
def generar_pdf():
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor('#1A202C')
    )
    story.append(Paragraph("<b>SOCIEDAD COOPERATIVA DE AHORRO Y PRÉSTAMO (NIVEL BÁSICO)</b>", title_style))
    story.append(Paragraph("<b>CÓMPUTO DE CAPITAL NETO Y REQUERIMIENTO DE CAPITALIZACIÓN — ANEXO U</b>", title_style))
    story.append(Spacer(1, 10))

    # Metadatos
    meta_data = [
        [Paragraph(f"<b>Sociedad:</b> {sociedad}", styles['Normal']), Paragraph(f"<b>Fecha de corte:</b> {fecha_corte.strftime('%d/%m/%Y')}", styles['Normal'])],
        [Paragraph(f"<b>Tipo de Cierre:</b> {tipo_cierre}", styles['Normal']), Paragraph("<b>Cifras en:</b> Miles de pesos mexicanos", styles['Normal'])]
    ]
    t_meta = Table(meta_data, colWidths=[340, 200])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F7FAFC')),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 12))

    # Cédula Oficial Anexo U
    renglones_tabla = [
        ["Renglón", "Concepto Regulatorio (Anexo U)", "Monto ($ Miles)"],
        ["(1)", "Cartera de crédito vigente", f"${c1:,.2f}"],
        ["(2)", "Cartera de crédito vencida", f"${c2:,.2f}"],
        ["(3)", "(-) Estimación preventiva para riesgos crediticios", f"${c3:,.2f}"],
        ["(4)", "Total de cartera de crédito neta [(1) + (2) - (3)]", f"${c4_cartera_neta:,.2f}"],
        ["(5)", "Requerimientos de capitalización [(4) * 8%]", f"${c5_req_capital:,.2f}"],
        ["(6)", "Capital Contable", f"${c6:,.2f}"],
        ["(7)", "(-) Certificados excedentes/voluntarios que no cumplen características", f"${c7:,.2f}"],
        ["(8)", "(-) Financiamiento para adquisición de partes sociales propias", f"${c8:,.2f}"],
        ["(9)", "Capital neto [(6) - (7) - (8)]", f"${c9_capital_neto:,.2f}"],
        ["NICAP", "Nivel de Capitalización [(9) / (5)] * 100", f"{nicap:,.2f}%"],
    ]

    t_cedula = Table(renglones_tabla, colWidths=[60, 360, 120])
    t_cedula.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2D3748')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9),
        ('ALIGN', (0,0), (0,-1), 'CENTER'),
        ('ALIGN', (2,0), (2,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0,4), (-1,4), colors.HexColor('#EDF2F7')),
        ('FONTNAME', (0,4), (-1,4), 'Helvetica-Bold'),
        ('BACKGROUND', (0,9), (-1,9), colors.HexColor('#EDF2F7')),
        ('FONTNAME', (0,9), (-1,9), 'Helvetica-Bold'),
        ('BACKGROUND', (0,10), (-1,10), colors.HexColor('#EBF8FF')),
        ('FONTNAME', (0,10), (-1,10), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('TOPPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_cedula)
    story.append(Spacer(1, 10))

    # Cuadro de Categoría Oficial
    color_map_pdf = {
        "Categoría A": colors.HexColor('#2E7D32'),
        "Categoría B": colors.HexColor('#1565C0'),
        "Categoría C": colors.HexColor('#E65100'),
        "Categoría D": colors.HexColor('#B71C1C'),
    }
    box_color = color_map_pdf.get(cat_nombre, colors.HexColor('#2D3748'))
    
    cat_data = [
        [Paragraph(f"<font color='white' size='11'><b>CATEGORÍA REGULATORIA DETERMINADA: {cat_nombre.upper()} ({cat_desc.upper()})</b></font>", ParagraphStyle('w', alignment=1))],
        [Paragraph(f"<font color='white' size='8'>Fundamento: Art. 15 LRASCAP | Cuadre Balanza: {'CUADRADO' if cuadra_balance else 'DESCUADRADO'} {(' | ' + cat_motivo_degradacion) if cat_motivo_degradacion else ''}</font>", ParagraphStyle('w2', alignment=1))]
    ]
    t_cat = Table(cat_data, colWidths=[540])
    t_cat.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), box_color),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_cat)
    story.append(Spacer(1, 10))

    # Resumen de Mejores Prácticas Nivel I
    story.append(Paragraph("<b>INDICADORES DE MEJORES PRÁCTICAS — NIVEL I (PRUDENCIAL)</b>", ParagraphStyle('Sub', fontName='Helvetica-Bold', fontSize=9, textColor=colors.HexColor('#0275D8'))))
    story.append(Spacer(1, 3))
    
    bp_data = [
        ["Indicador", "Cálculo Realizado", "Resultado", "Umbral / Parámetro"],
        ["Liquidez a Corto Plazo (≤30d)", f"${act_liq:,.2f} / ${pas_exig:,.2f}", f"{ratio_liq:.2f}%", "Mínimo regulatorio: ≥10%"],
        ["Cobertura Reserva Adjudicados", f"${res_adj:,.2f} / ${val_adj:,.2f}", f"{cob_adj:.2f}%", "Referencial Nivel Básico"],
        ["Activos en UDIS", f"{act_udis:,.2f} UDIS", f"{(act_udis/2500000.0)*100:.1f}%", "Límite régimen: 2,500,000 UDIS"]
    ]
    t_bp = Table(bp_data, colWidths=[150, 150, 90, 150])
    t_bp.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EBF8FF')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
        ('ALIGN', (2,0), (2,-1), 'CENTER'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('TOPPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_bp)

    doc.build(story)
    buffer.seek(0)
    return buffer

def generar_word():
    doc = Document()
    
    # Configuración de márgenes a 2 cm
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("SOCIEDAD COOPERATIVA DE AHORRO Y PRÉSTAMO (NIVEL BÁSICO)\nCÓMPUTO DE CAPITAL NETO Y REQUERIMIENTO DE CAPITALIZACIÓN — ANEXO U")
    title_run.bold = True
    title_run.font.size = Pt(12)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadatos
    doc.add_paragraph(f"Entidad: {sociedad} | Fecha de corte: {fecha_corte.strftime('%d/%m/%Y')} | Cierre: {tipo_cierre} (Cifras en miles de pesos)")

    # Tabla Cédula
    t = doc.add_table(rows=1, cols=3)
    t.autofit = True
    hdr = t.rows[0].cells
    hdr[0].text = "Renglón"
    hdr[1].text = "Concepto Regulatorio (Anexo U)"
    hdr[2].text = "Monto ($ Miles)"

    data = [
        ("(1)", "Cartera de crédito vigente", f"${c1:,.2f}"),
        ("(2)", "Cartera de crédito vencida", f"${c2:,.2f}"),
        ("(3)", "(-) Estimación preventiva para riesgos crediticios", f"${c3:,.2f}"),
        ("(4)", "Total de cartera de crédito neta [(1) + (2) - (3)]", f"${c4_cartera_neta:,.2f}"),
        ("(5)", "Requerimientos de capitalización [(4) * 8%]", f"${c5_req_capital:,.2f}"),
        ("(6)", "Capital Contable", f"${c6:,.2f}"),
        ("(7)", "(-) Certificados excedentes/voluntarios no computables", f"${c7:,.2f}"),
        ("(8)", "(-) Financiamiento para adquisición de partes sociales propias", f"${c8:,.2f}"),
        ("(9)", "Capital neto [(6) - (7) - (8)]", f"${c9_capital_neto:,.2f}"),
        ("NICAP", "Nivel de Capitalización [(9) / (5)] * 100", f"{nicap:,.2f}%"),
    ]

    for r, c, m in data:
        row = t.add_row().cells
        row[0].text = r
        row[1].text = c
        row[2].text = m

    doc.add_paragraph()
    
    # Cuadro Categoría
    p_cat = doc.add_paragraph()
    r_cat = p_cat.add_run(f"CLASIFICACIÓN RESULTANTE: {cat_nombre} ({cat_desc})\nNICAP Oficial: {nicap:,.2f}% | Cuadre Contable: {'Validado' if cuadra_balance else 'Descuadrado'}")
    r_cat.bold = True
    r_cat.font.size = Pt(11)
    if cat_motivo_degradacion:
        p_deg = doc.add_paragraph(f"Observación de Degradación: {cat_motivo_degradacion}")
        p_deg.runs[0].font.color.rgb = RGBColor(183, 28, 28)

    doc.add_paragraph("Indicadores de Mejores Prácticas Nivel I:")
    t_bp = doc.add_table(rows=1, cols=3)
    hdr_bp = t_bp.rows[0].cells
    hdr_bp[0].text = "Indicador"
    hdr_bp[1].text = "Resultado"
    hdr_bp[2].text = "Diagnóstico"
    
    bp_rows = [
        ("Liquidez Corto Plazo (≤30d)", f"{ratio_liq:.2f}%", "Solvente" if ratio_liq >= 10.0 else "Riesgo de liquidez"),
        ("Cobertura Bienes Adjudicados", f"{cob_adj:.2f}%", "Evaluación referencial de cobertura"),
        ("Consumo Límite UDIS", f"{act_udis:,.2f} UDIS", "Bajo Nivel Básico" if act_udis < 2300000 else ("Alerta Límite Nivel Básico" if act_udis < 2500000 else "Superó Nivel Básico"))
    ]
    for ind, res, diag in bp_rows:
        row = t_bp.add_row().cells
        row[0].text = ind
        row[1].text = res
        row[2].text = diag

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==========================================
# BOTONES DE EXPORTACIÓN
# ==========================================
st.markdown("---")
st.subheader("📑 Descarga de Informes Oficiales")
st.caption("Genera los expedientes listos para firmas del Consejo y archivo de auditoría:")

col_btn1, col_btn2 = st.columns(2)
with col_btn1:
    pdf_bytes = generar_pdf()
    st.download_button(
        label="📄 Descargar Cédula en PDF (1 Página)",
        data=pdf_bytes,
        file_name=f"Anexo_U_NICAP_{fecha_corte.strftime('%Y%m%d')}.pdf",
        mime="application/pdf",
        use_container_width=True
    )

with col_btn2:
    docx_bytes = generar_word()
    st.download_button(
        label="📝 Descargar Cédula en Word (.docx)",
        data=docx_bytes,
        file_name=f"Anexo_U_NICAP_{fecha_corte.strftime('%Y%m%d')}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )